# -*- coding: utf-8 -*-
# ------------------------------------------------------------------
# 听课搭子 (tingke-dazi) · Copyright (c) 2026 luoxiu065-zjx
# 原始仓库:https://github.com/luoxiu065-zjx/tingke-dazi
# 转载、二次发布请保留本版权声明和原作者署名。build: tkdz-1bf2795e
# ------------------------------------------------------------------
"""课件模块:pdf / pptx / md / txt 抽成按页文本,存 课件/<课号>/;上课时按最近转写的关键词挑最相关的几页喂给总结和提问。
顺带把课件里的大写术语补进 术语/<课号>.txt。不 import app;tests/test_slides.py 测。"""
import os, re, json, datetime

DIR = None
TERMS_DIR = None
STOP = set("""the a an and or of to in on for with by from as at is are was were be been this that these those it its we you they he she
i our your their his her not no yes but if then so than too very can could would should will may might do does did have has had
what which who whom whose where when why how all any some more most other such only own same just also about into over under
between after before during while because until again further here there each few both very s t re ve ll d m""".split())


def init(base):
    global DIR, TERMS_DIR
    DIR = os.path.join(base, "课件")
    TERMS_DIR = os.path.join(base, "术语")
    os.makedirs(DIR, exist_ok=True)


def _code_dir(code):
    d = os.path.join(DIR, re.sub(r'[\\/:*?"<>|]', "_", code or "其他"))
    os.makedirs(d, exist_ok=True)
    return d


# ---------- 抽文本 ----------
def extract_pages(path):
    """返回 [{"n": 页码, "text": 文本}],空页不要。"""
    ext = os.path.splitext(path)[1].lower()
    pages = []
    if ext == ".pdf":
        import fitz
        with fitz.open(path) as doc:
            for i, pg in enumerate(doc, 1):
                t = pg.get_text("text")
                if t and t.strip():
                    pages.append({"n": i, "text": _tidy(t)})
    elif ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(path)
        for i, slide in enumerate(prs.slides, 1):
            parts = []
            for shp in slide.shapes:
                if shp.has_text_frame:
                    parts.append(shp.text_frame.text)
                if getattr(shp, "has_table", False) and shp.has_table:
                    for row in shp.table.rows:
                        parts.append(" | ".join(c.text for c in row.cells))
            try:
                if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                    parts.append("[讲者备注] " + slide.notes_slide.notes_text_frame.text)
            except Exception:
                pass
            t = "\n".join(p for p in parts if p and p.strip())
            if t.strip():
                pages.append({"n": i, "text": _tidy(t)})
    elif ext in (".md", ".txt"):
        with open(path, encoding="utf-8", errors="replace") as f:
            txt = f.read()
        chunks = re.split(r"\n(?=#{1,3} )", txt) if "\n#" in txt else [txt[i:i + 1500] for i in range(0, len(txt), 1500)]
        for i, c in enumerate(chunks, 1):
            if c.strip():
                pages.append({"n": i, "text": _tidy(c)})
    else:
        raise ValueError("不支持的文件类型:%s(只收 pdf / pptx / md / txt)" % ext)
    return pages


def _tidy(t):
    t = t.replace("\x00", " ")
    t = re.sub(r"[ \t]+", " ", t)
    t = re.sub(r"\n{3,}", "\n\n", t)
    return t.strip()


# ---------- 存取 ----------
def add(code, src_path, display_name=None):
    """抽文本存成 课件/<课号>/<名>.json,返回摘要。同名覆盖。"""
    name = display_name or os.path.basename(src_path)
    pages = extract_pages(src_path)
    if not pages:
        raise ValueError("这份课件里抽不出文字(可能是扫描图片)")
    rec = {"name": name, "added": datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), "pages": pages,
           "chars": sum(len(p["text"]) for p in pages)}
    with open(os.path.join(_code_dir(code), _safe(name) + ".json"), "w", encoding="utf-8") as f:
        json.dump(rec, f, ensure_ascii=False)
    added_terms = add_terms(code, pages)
    return {"name": name, "pages": len(pages), "chars": rec["chars"], "new_terms": added_terms}


def _safe(name):
    return re.sub(r'[\\/:*?"<>|]', "_", name)[:80]


def list_files(code):
    out = []
    for fn in sorted(os.listdir(_code_dir(code))):
        if fn.endswith(".json"):
            try:
                with open(os.path.join(_code_dir(code), fn), encoding="utf-8") as f:
                    r = json.load(f)
                out.append({"name": r["name"], "added": r.get("added", ""), "pages": len(r["pages"]), "chars": r.get("chars", 0)})
            except Exception:
                pass
    return out


def remove(code, name):
    p = os.path.join(_code_dir(code), _safe(name) + ".json")
    if os.path.exists(p):
        os.remove(p); return True
    return False


def _all_pages(code):
    for fn in os.listdir(_code_dir(code)):
        if fn.endswith(".json"):
            try:
                with open(os.path.join(_code_dir(code), fn), encoding="utf-8") as f:
                    r = json.load(f)
                for p in r["pages"]:
                    yield r["name"], p
            except Exception:
                pass


# ---------- 检索 ----------
def _tokens(text):
    return [w for w in re.findall(r"[A-Za-z][A-Za-z\-']{2,}", text.lower()) if w not in STOP]


def select(code, query_text, limit=1500, top=3):
    """按最近转写的关键词挑最相关的几页,拼成 ≤limit 字的节选;没课件或没匹配返回空串。"""
    q = _tokens(query_text or "")
    if not q:
        return ""
    qset = {}
    for w in q:
        qset[w] = qset.get(w, 0) + 1
    scored = []
    for name, p in _all_pages(code):
        toks = set(_tokens(p["text"]))
        if not toks:
            continue
        score = sum(c for w, c in qset.items() if w in toks)
        if score:
            scored.append((score / (1 + len(toks) ** 0.5), name, p))
    scored.sort(key=lambda x: -x[0])
    out, used = [], 0
    for _, name, p in scored[:top]:
        piece = "《%s》第 %d 页:%s" % (name, p["n"], p["text"][:600])
        if used + len(piece) > limit:
            piece = piece[:max(0, limit - used)]
        if not piece:
            break
        out.append(piece); used += len(piece)
    return "\n---\n".join(out)


# ---------- 术语补全 ----------
def extract_terms(pages, max_terms=60):
    """课件里反复出现的大写词/缩写/专有名词(出现 ≥2 次),当作术语候选。"""
    cnt = {}
    for p in pages:
        for w in re.findall(r"\b(?:[A-Z][a-z]{3,}(?: [A-Z][a-z]{3,}){0,2}|[A-Z]{2,}[A-Za-z0-9\-]*)\b", p["text"]):
            if w.lower() in STOP or len(w) < 3:
                continue
            cnt[w] = cnt.get(w, 0) + 1
    terms = [w for w, c in sorted(cnt.items(), key=lambda x: -x[1]) if c >= 2]
    return terms[:max_terms]


def add_terms(code, pages):
    """把新术语追加进 术语/<课号>.txt,返回新增的词。"""
    cand = extract_terms(pages)
    if not cand or not TERMS_DIR:
        return []
    os.makedirs(TERMS_DIR, exist_ok=True)
    p = os.path.join(TERMS_DIR, code + ".txt")
    have = set()
    if os.path.exists(p):
        with open(p, encoding="utf-8") as f:
            have = {w.strip().lower() for w in re.split(r"[,\n]", f.read()) if w.strip()}
    new = [w for w in cand if w.lower() not in have]
    if new:
        with open(p, "a", encoding="utf-8") as f:
            f.write("\n# 课件自动补充 %s\n%s\n" % (datetime.date.today().isoformat(), ", ".join(new)))
    return new
